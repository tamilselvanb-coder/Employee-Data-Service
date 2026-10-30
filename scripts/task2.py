import csv
import json
import logging
from pathlib import Path

import openpyxl

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).resolve().parent.parent


def process_files() -> None:
    csv_path = BASE_DIR / "employees.csv"
    json_path = BASE_DIR / "employees.json"
    high_earners_path = BASE_DIR / "high_earners.csv"

    employees: list[dict] = []
    high_earners: list[dict] = []

    with open(csv_path, mode="r") as f:
        reader = csv.DictReader(f)
        for row in reader:
            cleaned_row = {
                "name": row["name"].strip(),
                "department": row["department"].strip(),
                "salary": int(row["salary"].strip()),
            }
            employees.append(cleaned_row)
            if cleaned_row["salary"] > 50000:
                high_earners.append(cleaned_row)

    with open(json_path, mode="w") as f:
        json.dump(employees, f, indent=4)

    with open(high_earners_path, mode="w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["name", "department", "salary"])
        writer.writeheader()
        writer.writerows(high_earners)

    print(f"Rows read: {len(employees)} High earners written: {len(high_earners)}")


def excel_operations() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Salaries"

    csv_path = BASE_DIR / "employees.csv"
    excel_path = BASE_DIR / "salary.xlsx"

    with open(csv_path, "r") as f:
        reader = csv.reader(f)
        for row in reader:
            ws.append([r.strip() for r in row])

    wb.save(excel_path)

    wb_read = openpyxl.load_workbook(excel_path)
    sheet = wb_read.active
    print(f"salary.xlsx created and read back: {sheet.max_row} rows")


def read_file(path: str | Path) -> dict | list[dict] | None:
    file_path = BASE_DIR / path if isinstance(path, str) else path
    try:
        with open(file_path, "r") as f:
            return json.load(f)
    except FileNotFoundError:
        logger.error(f"File not found: {path}")
    except json.JSONDecodeError:
        logger.error(f"JSON decode error in file: {path}")
    return None


def find_employee(name: str) -> dict | None:
    employees = read_file("employees.json")
    if isinstance(employees, list):
        for emp in employees:
            if emp.get("name", "").lower() == name.lower():
                return emp
    return None


if __name__ == "__main__":
    process_files()
    excel_operations()
    read_file("missing.txt")
